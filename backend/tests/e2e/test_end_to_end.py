"""Upload to cited answer through the real pipeline.

Runs only with ``PARSRAG_E2E=1`` and a reachable Qdrant (``QDRANT_HOST`` /
``QDRANT_PORT``). Parsing, Persian normalization, chunking, the e5
embedding model, Qdrant storage and search, reranking, prompt building, and
SSE streaming are all real. Only the language model is a deterministic stand-in
that cites the first excerpt it is given, so the test checks what retrieval put
first without needing a GPU or an API key.
"""

from __future__ import annotations

import io
import json
import os
import re
import uuid
from collections.abc import Generator, Iterator
from typing import Any

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from llama_index.core import Settings
from llama_index.core.llms import (
    CompletionResponse,
    CompletionResponseGen,
    CustomLLM,
    LLMMetadata,
)
from llama_index.core.llms.callbacks import llm_completion_callback

pytestmark = pytest.mark.skipif(
    os.getenv("PARSRAG_E2E") != "1", reason="set PARSRAG_E2E=1 with Qdrant running"
)

_FIRST_EXCERPT = re.compile(r"^\[1\][^\n]*\n(?P<text>[^\n]+)", re.MULTILINE)

# Eight unrelated topics; each question paraphrases its section.
SECTIONS = {
    "فتوسنتز": (
        "گیاهان سبز با کمک سبزینه نور خورشید را جذب می‌کنند و از آب و کربن دی‌اکسید "
        "قند می‌سازند. اکسیژن هوا محصول جانبی همین فرایند است."
    ),
    "زمین‌لرزه": (
        "حرکت ناگهانی صفحه‌های پوستهٔ زمین در امتداد گسل‌ها انرژی ذخیره‌شده را آزاد می‌کند. "
        "بزرگی این رویداد با مقیاس ریشتر سنجیده می‌شود."
    ),
    "قانون اهم": (
        "در یک رسانا، شدت جریان با ولتاژ دو سر آن نسبت مستقیم و با مقاومت نسبت وارون دارد؛ "
        "یعنی ولتاژ برابر حاصل‌ضرب جریان در مقاومت است."
    ),
    "تخت جمشید": (
        "کاخ‌های هخامنشی در نزدیکی شیراز به فرمان داریوش بزرگ ساخته شدند. ستون‌های سنگی "
        "و نقش برجستهٔ نمایندگان ملت‌ها از نشانه‌های این بناست."
    ),
    "دیابت نوع دو": (
        "در این بیماری بدن به انسولین پاسخ مناسب نمی‌دهد و قند خون بالا می‌ماند. ورزش "
        "منظم و رژیم غذایی کم‌قند در کنترل آن مؤثرند."
    ),
    "رمزنگاری RSA": (
        "امنیت این روش کلید عمومی بر دشواری تجزیهٔ حاصل‌ضرب دو عدد اول بزرگ تکیه دارد. "
        "کلید عمومی برای رمز کردن و کلید خصوصی برای گشودن پیام به کار می‌رود."
    ),
    "چرخهٔ آب": (
        "آب اقیانوس‌ها با گرمای خورشید بخار می‌شود، در هوا سرد و به ابر تبدیل می‌شود و "
        "سپس به صورت باران یا برف به زمین بازمی‌گردد."
    ),
    "بیت‌کوین": (
        "این ارز دیجیتال بدون بانک مرکزی کار می‌کند و تراکنش‌ها در دفترکلی توزیع‌شده به نام "
        "زنجیرهٔ بلوک ثبت می‌شوند که استخراج‌کنندگان آن را تأیید می‌کنند."
    ),
}
QUESTIONS = {
    "فتوسنتز": "گیاهان چطور از نور آفتاب غذا درست می‌کنند؟",
    "زمین‌لرزه": "شدت لرزش زمین با چه معیاری اندازه گرفته می‌شود؟",
    "قانون اهم": "رابطهٔ ولتاژ و جریان الکتریکی در یک سیم چیست؟",
    "تخت جمشید": "کدام پادشاه دستور ساخت کاخ‌های نزدیک شیراز را داد؟",
    "دیابت نوع دو": "برای مهار بالا بودن قند خون چه پیشنهاد شده است؟",
    "رمزنگاری RSA": "امنیت روش کلید عمومی بر چه مسئله‌ای استوار است؟",
    "چرخهٔ آب": "باران چگونه از دریا به خشکی می‌رسد؟",
    "بیت‌کوین": "تراکنش‌های این پول دیجیتال کجا ثبت می‌شوند؟",
}


class CitingLLM(CustomLLM):
    """Answer with the first context excerpt and cite it as [1]."""

    @property
    def metadata(self) -> LLMMetadata:
        """Describe a small local model."""
        return LLMMetadata(context_window=8192, num_output=256, model_name="citing")

    @staticmethod
    def answer(prompt: str) -> str:
        """Quote the start of excerpt [1], or refuse when there is none."""
        match = _FIRST_EXCERPT.search(prompt)
        if match is None:
            return "بر اساس اسناد ارائه شده، پاسخی برای این سوال در متن یافت نشد."
        return f"بر پایهٔ سند: {match.group('text')[:120]} [1]"

    @llm_completion_callback()
    def complete(
        self, prompt: str, formatted: bool = False, **kwargs: Any
    ) -> CompletionResponse:
        """Return the whole answer."""
        return CompletionResponse(text=self.answer(prompt))

    @llm_completion_callback()
    def stream_complete(
        self, prompt: str, formatted: bool = False, **kwargs: Any
    ) -> CompletionResponseGen:
        """Yield the answer word by word."""

        def generate() -> Generator[CompletionResponse, None, None]:
            text = ""
            for word in self.answer(prompt).split(" "):
                delta = word + " "
                text += delta
                yield CompletionResponse(text=text, delta=delta)

        return generate()


@pytest.fixture(scope="module")
def client() -> Iterator[TestClient]:
    """Run the app on a throwaway Qdrant collection with real embeddings."""
    collection = f"parsrag_e2e_{uuid.uuid4().hex[:10]}"
    os.environ["QDRANT_COLLECTION"] = collection
    os.environ["KEYWORD_SEARCH"] = "1"
    from backend.api import dependencies
    from backend.infrastructure.llm import factory

    factory.setup_llm_and_embeddings()
    Settings.llm = CitingLLM()
    factory._condense_llm = CitingLLM()
    dependencies._shared_document_repository.cache_clear()
    from backend.main import app

    with TestClient(app) as test_client:
        yield test_client
    repository = dependencies._shared_document_repository()
    repository.client.delete_collection(collection)  # type: ignore[attr-defined]
    dependencies._shared_document_repository.cache_clear()
    os.environ.pop("QDRANT_COLLECTION", None)


def _persian_pdf() -> bytes:
    """Two text pages: an overview and the fact the question asks about."""
    document = pymupdf.open()
    pages = [
        "پارس‌رگ اسناد فارسی را نمایه می‌کند و به پرسش‌ها با ارجاع پاسخ می‌دهد.",
        "در نسخهٔ نهایی، بودجهٔ پروژه دویست میلیون تومان و مدت اجرای آن شش ماه تعیین شد.",
    ]
    for text in pages:
        page = document.new_page()
        page.insert_htmlbox(pymupdf.Rect(50, 50, 545, 800), f'<p dir="rtl">{text}</p>')
    return bytes(document.tobytes())


def _topics_docx() -> bytes:
    document = Document()
    for title, body in SECTIONS.items():
        document.add_heading(title, level=1)
        document.add_paragraph(body)
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


def _events(body: str) -> list[tuple[str, Any]]:
    events = []
    for block in body.strip().split("\n\n"):
        fields = dict(line.split(": ", 1) for line in block.splitlines())
        events.append((fields["event"], json.loads(fields["data"])))
    return events


def test_persian_pdf_upload_streams_a_cited_answer_from_the_right_page(
    client: TestClient,
) -> None:
    session = f"e2e-{uuid.uuid4().hex[:8]}"
    upload = client.post(
        "/ingest",
        files={"files": ("budget.pdf", _persian_pdf(), "application/pdf")},
        data={"session_id": session},
    )
    assert upload.status_code == 200, upload.text
    assert upload.json()["files"][0]["chunks"] >= 2

    answer = client.post(
        "/query/stream",
        json={
            "prompt": "هزینهٔ پروژه چقدر در نظر گرفته شد؟",
            "mode": "strict",
            "session_id": session,
        },
    )
    assert answer.status_code == 200
    events = _events(answer.text)
    stages = [data["stage"] for kind, data in events if kind == "stage"]
    assert stages == ["understanding", "retrieving", "generating"]
    kind, done = events[-1]
    assert kind == "done"
    assert done["cited"] == [1]
    first = done["source_nodes"][0]
    assert first["metadata"]["filename"] == "budget.pdf"
    # Page 2 itself, or the bridge chunk that spans pages 1-2.
    assert first["metadata"].get("page_end", first["metadata"]["page"]) == 2
    # Presentation-form glyphs in the PDF text layer are normalized for search.
    assert "بودجه" in first["text"]

    deleted = client.delete(f"/sessions/{session}")
    assert deleted.status_code == 200
    assert client.get(f"/sessions/{session}/files").json() == []


def test_retrieval_puts_the_answering_section_first(client: TestClient) -> None:
    """Quality gate: paraphrased Persian questions find their own section."""
    session = f"e2e-{uuid.uuid4().hex[:8]}"
    upload = client.post(
        "/ingest",
        files={"files": ("topics.docx", _topics_docx(), "application/octet-stream")},
        data={"session_id": session},
    )
    assert upload.status_code == 200, upload.text

    hits = []
    for topic, question in QUESTIONS.items():
        response = client.post(
            "/query", json={"prompt": question, "mode": "hybrid", "session_id": session}
        )
        assert response.status_code == 200, response.text
        result = response.json()
        cited = result["source_nodes"][result["cited"][0] - 1]["text"]
        hits.append(topic in cited)

    client.delete(f"/sessions/{session}")
    assert sum(hits) >= 7, (
        f"only {sum(hits)} of {len(hits)} questions found their section"
    )
