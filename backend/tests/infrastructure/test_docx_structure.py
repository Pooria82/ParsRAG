"""Word documents are chunked by heading instead of one chunk per line."""

from collections.abc import Callable
from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument

from backend.core.domain.documents import ExtractedNode
from backend.core.strategies.multi_doc_utils import format_multi_doc_context
from backend.infrastructure.parsers.document_parser import parse_document_sections


def _docx(build: Callable[[DocxDocument], None]) -> bytes:
    document = Document()
    build(document)
    payload = BytesIO()
    document.save(payload)
    return payload.getvalue()


def test_paragraphs_under_a_heading_form_one_section_with_its_path() -> None:
    def build(document: DocxDocument) -> None:
        document.add_heading("گزارش فاز اول", level=0)
        document.add_heading("ابزارهای تست", level=1)
        document.add_paragraph("Katalon Recorder برای تست رابط کاربری استفاده شد.")
        document.add_paragraph("JMeter بار سرور را اندازه گرفت.")
        document.add_heading("نتیجه", level=1)
        document.add_paragraph("همه آزمون‌ها موفق بودند.")

    sections = parse_document_sections(_docx(build), "report.docx")

    assert [section.text for section in sections] == [
        "گزارش فاز اول › ابزارهای تست\n"
        "Katalon Recorder برای تست رابط کاربری استفاده شد.\n"
        "JMeter بار سرور را اندازه گرفت.",
        "گزارش فاز اول › نتیجه\nهمه آزمون‌ها موفق بودند.",
    ]
    assert sections[0].metadata == {"paragraph": 3, "paragraph_end": 4}
    assert sections[1].metadata == {"paragraph": 6}


def test_long_sections_split_near_one_chunk_and_keep_the_heading() -> None:
    sentence = " ".join(["واژه"] * 100)

    def build(document: DocxDocument) -> None:
        document.add_heading("فصل دوم", level=1)
        for _ in range(3):
            document.add_paragraph(sentence)

    sections = parse_document_sections(_docx(build), "long.docx")

    assert len(sections) == 3
    assert all(section.text.startswith("فصل دوم\n") for section in sections)
    assert [section.metadata["paragraph"] for section in sections] == [2, 3, 4]


def test_tables_carry_their_heading_path() -> None:
    def build(document: DocxDocument) -> None:
        document.add_heading("هزینه‌ها", level=2)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "مورد", "مبلغ"
        table.cell(1, 0).text, table.cell(1, 1).text = "سرور", "۲۰۰"

    sections = parse_document_sections(_docx(build), "costs.docx")

    assert sections[0].text.startswith("هزینه‌ها\n[جدول 1]\n| مورد | مبلغ |")
    assert sections[0].metadata == {"paragraph": 1, "table": 1}


def test_one_cell_boxes_stay_with_the_sentence_that_introduces_them() -> None:
    """A command in a one-cell box is indexed with the text that leads to it."""

    def build(document: DocxDocument) -> None:
        document.add_paragraph("مدل محلی را با دستور زیر اجرا کنید:")
        box = document.add_table(rows=1, cols=1)
        box.cell(0, 0).text = "ollama run qwen:2.5-coder-1.5b"
        document.add_paragraph("سپس افزونه را باز کنید.")

    sections = parse_document_sections(_docx(build), "setup.docx")

    assert len(sections) == 1
    assert sections[0].text == (
        "مدل محلی را با دستور زیر اجرا کنید:\nollama run qwen:2.5-coder-1.5b\n"
        "سپس افزونه را باز کنید."
    )
    assert sections[0].metadata == {"paragraph": 1, "paragraph_end": 2}


def test_tables_are_cited_by_number_caption_and_preceding_paragraph() -> None:
    """A table is cited as "paragraph N · table M", with its caption above it."""

    def build(document: DocxDocument) -> None:
        document.add_paragraph("مقدمه")
        document.add_paragraph("جدول ۳-۶: ابزارهای متن‌باز")
        first = document.add_table(rows=1, cols=2)
        first.cell(0, 0).text, first.cell(0, 1).text = "PyAutoGUI", "10.9k"
        document.add_paragraph("توضیح میانی")
        second = document.add_table(rows=1, cols=2)
        second.cell(0, 0).text, second.cell(0, 1).text = "بدون شرح", "۲"

    sections = parse_document_sections(_docx(build), "tools.docx")

    tables = [s for s in sections if "table" in s.metadata]
    assert [s.metadata for s in tables] == [
        {"paragraph": 2, "table": 1},
        {"paragraph": 3, "table": 2},
    ]
    assert tables[0].text.startswith(
        "[جدول 1] جدول ۳-۶: ابزارهای متن‌باز\n| PyAutoGUI |"
    )
    assert tables[1].text.startswith("[جدول 2]\n") and "بدون شرح" in tables[1].text


def test_documents_without_headings_still_group_short_lines() -> None:
    def build(document: DocxDocument) -> None:
        for index in range(5):
            document.add_paragraph(f"خط کوتاه شماره {index}")

    sections = parse_document_sections(_docx(build), "notes.docx")

    assert len(sections) == 1
    assert sections[0].metadata == {"paragraph": 1, "paragraph_end": 5}


def test_paragraph_ranges_are_cited_as_ranges() -> None:
    context = format_multi_doc_context(
        [
            ExtractedNode(
                text="متن",
                metadata={"filename": "r.docx", "paragraph": 3, "paragraph_end": 7},
            )
        ]
    )
    assert "[1] r.docx, paragraphs: 3-7\nمتن" in context
