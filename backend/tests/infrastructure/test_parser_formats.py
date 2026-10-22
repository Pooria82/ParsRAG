"""Parser coverage for grouped shapes, notes, content controls, tables, and limits."""

import time
from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend.infrastructure.parsers.document_parser import (
    _decode_text,
    _pptx_pictures,
    parse_document_file,
    recognize_images,
)
from backend.infrastructure.parsers.ocr import OCRError, OCRSettings, extract_image_text


def _png(size: tuple[int, int] = (400, 200)) -> bytes:
    output = BytesIO()
    Image.new("RGB", size, "white").save(output, format="PNG")
    return output.getvalue()


def _pptx(build: object) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    build(slide)  # type: ignore[operator]
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_text_and_pictures_inside_grouped_shapes_are_indexed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("OCR_ENABLED", "1")

    def build(slide: object) -> None:
        group = slide.shapes.add_group_shape()  # type: ignore[attr-defined]
        group.shapes.add_textbox(
            Inches(1), Inches(1), Inches(4), Inches(1)
        ).text = "Grouped architecture note"
        group.shapes.add_picture(BytesIO(_png()), Inches(1), Inches(2))

    with patch(
        "backend.infrastructure.parsers.document_parser.extract_image_text",
        return_value="text inside a grouped picture",
    ):
        parsed = parse_document_file(_pptx(build), "deck.pptx")

    texts = [section.text for section in parsed.sections]
    assert any("Grouped architecture note" in text for text in texts)
    assert any("text inside a grouped picture" in text for text in texts)
    assert {"slide": 1, "image": 1} in [section.metadata for section in parsed.sections]


def test_speaker_notes_are_part_of_the_slide(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("OCR_ENABLED", "0")

    def build(slide: object) -> None:
        slide.shapes.add_textbox(  # type: ignore[attr-defined]
            Inches(1), Inches(1), Inches(4), Inches(1)
        ).text = "Release plan"
        slide.notes_slide.notes_text_frame.text = (  # type: ignore[attr-defined]
            "Explain the rollback steps here"
        )

    parsed = parse_document_file(_pptx(build), "deck.pptx")

    assert parsed.sections[0].metadata == {"slide": 1}
    assert "Explain the rollback steps here" in parsed.sections[0].text


def test_a_linked_picture_adds_a_notice_instead_of_rejecting_the_deck() -> None:
    linked = MagicMock()
    linked.has_table = False
    linked.shape_type = 13  # MSO_SHAPE_TYPE.PICTURE
    type(linked).image = property(lambda self: (_ for _ in ()).throw(ValueError()))
    embedded = MagicMock(has_table=False, shape_type=13)
    embedded.image.blob = b"picture"
    slide = MagicMock()
    slide.shapes = [linked, embedded]
    presentation = MagicMock()
    presentation.slides = [slide]
    notices: list[str] = []

    assert _pptx_pictures(presentation, notices) == [b"", b"picture"]
    assert notices == ["ocr_partial"]


def test_content_inside_word_content_controls_is_indexed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("OCR_ENABLED", "0")
    document = Document()
    document.add_paragraph("Visible introduction")
    control = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r>"
        "<w:t>Form field answer: 1404</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )
    document.element.body.insert(1, control)
    output = BytesIO()
    document.save(output)

    parsed = parse_document_file(output.getvalue(), "form.docx")

    assert "Form field answer: 1404" in " ".join(s.text for s in parsed.sections)


def test_images_after_the_document_deadline_are_skipped_with_a_notice() -> None:
    settings = OCRSettings(True, "fas+eng", 200, 10, 5, deadline=time.monotonic() - 1)
    notices: list[str] = []
    with patch(
        "backend.infrastructure.parsers.document_parser.extract_image_text"
    ) as extract:
        texts = recognize_images([b"a", b"b"], settings, notices)

    assert texts == ["", ""]
    assert notices == ["ocr_partial", "ocr_partial"]
    extract.assert_not_called()


def test_pdf_pages_after_the_deadline_keep_native_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("OCR_ENABLED", "1")
    monkeypatch.setenv("OCR_DOCUMENT_TIMEOUT_SECONDS", "30")
    page = MagicMock()
    page.get_text.return_value = ""
    document = MagicMock()
    document.__iter__.return_value = [page]
    with (
        patch(
            "backend.infrastructure.parsers.document_parser.pymupdf.open",
            return_value=document,
        ),
        patch(
            "backend.infrastructure.parsers.document_parser.render_page_png",
            return_value=b"png",
        ),
        patch(
            "backend.infrastructure.parsers.document_parser.recognize_page_png"
        ) as recognize,
        patch(
            "backend.infrastructure.parsers.ocr.time.monotonic",
            side_effect=[0.0, 100.0, 100.0],
        ),
        pytest.raises(Exception, match="readable text"),
    ):
        parse_document_file(b"pdf", "scan.pdf")

    recognize.assert_not_called()


def test_scanned_pdf_pages_are_recognized_in_parallel_in_page_order(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("OCR_ENABLED", "1")
    monkeypatch.setenv("OCR_WORKERS", "4")
    pages = [MagicMock() for _ in range(4)]
    for page in pages:
        page.get_text.return_value = ""
    document = MagicMock()
    document.__iter__.return_value = pages
    rendered = iter([b"p1", b"p2", b"p3", b"p4"])
    with (
        patch(
            "backend.infrastructure.parsers.document_parser.pymupdf.open",
            return_value=document,
        ),
        patch(
            "backend.infrastructure.parsers.document_parser.render_page_png",
            side_effect=lambda page, settings: next(rendered),
        ),
        patch(
            "backend.infrastructure.parsers.document_parser.recognize_page_png",
            side_effect=lambda image, settings: f"text of {image.decode()}",
        ),
    ):
        parsed = parse_document_file(b"pdf", "scan.pdf")

    assert [(s.text, s.metadata) for s in parsed.sections] == [
        (f"text of p{n}", {"page": n}) for n in range(1, 5)
    ]


def test_a_decompression_bomb_is_an_ocr_error_not_a_crash(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 1000)
    with pytest.raises(OCRError, match="pixel limit"):
        extract_image_text(_png((200, 200)), OCRSettings(True, "fas+eng", 200, 10, 5))


def test_html_table_cells_are_separated() -> None:
    parsed = parse_document_file(
        b"<table><tr><th>Name</th><th>City</th></tr>"
        b"<tr><td>Ali</td><td>Tabriz</td></tr></table>",
        "people.html",
    )
    text = parsed.sections[0].text
    assert "NameCity" not in text and "AliTabriz" not in text
    assert "Ali | Tabriz" in text


def test_a_stray_byte_in_persian_utf8_keeps_the_persian_text() -> None:
    persian = "این گزارش آزمون نرم‌افزار است و نتیجه‌ها را نشان می‌دهد".encode()
    decoded = _decode_text(persian + b" caf\xe9")
    assert decoded.startswith("این گزارش آزمون")


def test_a_western_legacy_file_still_decodes_as_windows_1252() -> None:
    assert _decode_text("café crème".encode("cp1252")) == "café crème"


@pytest.mark.parametrize(
    ("name", "separator"), [("people.csv", ","), ("people.tsv", "\t")]
)
def test_delimited_files_become_tables_with_row_records(
    name: str, separator: str
) -> None:
    data = separator.join(["name", "city"]) + "\n" + separator.join(["Ali", "Tabriz"])
    parsed = parse_document_file(data.encode(), name)

    assert parsed.sections[0].metadata == {"table": 1}
    assert "Ali" in parsed.sections[0].text
    assert "[city]: Tabriz" in parsed.sections[0].text


def test_markdown_is_split_at_headings_and_long_text_by_size() -> None:
    markdown = b"# Intro\nfirst part\n\n## Setup\nsecond part\n"
    sections = parse_document_file(markdown, "guide.md").sections
    assert [s.metadata for s in sections] == [{"section": 1}, {"section": 2}]
    assert sections[1].text.startswith("## Setup")

    long_text = ("word " * 50 + "\n") * 20
    sections = parse_document_file(long_text.encode(), "notes.txt").sections
    assert len(sections) == 4
    assert [s.metadata["section"] for s in sections] == [1, 2, 3, 4]


def test_an_uploaded_image_keeps_a_preview(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("OCR_ENABLED", "1")
    with patch(
        "backend.infrastructure.parsers.document_parser.extract_image_text",
        return_value="screenshot text",
    ):
        parsed = parse_document_file(_png(), "screen.png")

    assert parsed.sections[0].metadata == {"page": 1, "image": 1}
    assert parsed.images[1][:4] == b"RIFF"
