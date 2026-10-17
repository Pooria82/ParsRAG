import json
from io import BytesIO
from unittest.mock import MagicMock, patch

import pymupdf
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from backend.core.domain.documents import ParsedSection
from backend.core.domain.exceptions import DocumentError, EmptyDocumentError
from backend.infrastructure.parsers.chunker import bridge_adjacent_pages, chunk_text
from backend.infrastructure.parsers.document_parser import (
    parse_document,
    parse_document_file,
    parse_document_sections,
)


@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_parse_empty_pdf(mock_fitz_open: MagicMock) -> None:
    # Mock a PDF document with one page but no text
    mock_doc = MagicMock()
    mock_page = MagicMock()
    mock_page.get_text.return_value = ""

    # Doc behaves like an iterator of pages
    mock_doc.__iter__.return_value = [mock_page]
    mock_fitz_open.return_value = mock_doc

    with pytest.raises(EmptyDocumentError, match="require OCR") as raised:
        parse_document(b"fake pdf bytes", "fake.pdf")
    assert raised.value.code == "ocr_disabled"

    mock_doc.close.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_pdf_sections_keep_page_numbers(mock_fitz_open: MagicMock) -> None:
    mock_doc = MagicMock()
    first, second = MagicMock(), MagicMock()
    first.get_text.return_value = "first page"
    second.get_text.return_value = "second page"
    mock_doc.__iter__.return_value = [first, second]
    mock_fitz_open.return_value = mock_doc

    sections = parse_document_sections(b"pdf", "guide.pdf")

    assert [section.metadata for section in sections] == [{"page": 1}, {"page": 2}]
    mock_doc.close.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.extract_page_text")
@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_scanned_pdf_uses_ocr_when_enabled(
    mock_fitz_open: MagicMock,
    mock_extract_page_text: MagicMock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Fully scanned PDFs use OCR and preserve their page citation."""
    monkeypatch.setenv("OCR_ENABLED", "1")
    page = MagicMock()
    page.get_text.return_value = ""
    document = MagicMock()
    document.__iter__.return_value = [page]
    mock_fitz_open.return_value = document
    mock_extract_page_text.return_value = "متن فارسی اسکن‌شده"

    sections = parse_document_sections(b"pdf", "scan.pdf")

    assert sections[0].text == "متن فارسی اسکن‌شده"
    assert sections[0].metadata == {"page": 1}
    mock_extract_page_text.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.extract_page_text")
@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_garbled_text_layers_are_replaced_by_ocr(
    mock_fitz_open: MagicMock,
    mock_extract_page_text: MagicMock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Pages whose fonts lack Unicode maps are read by OCR, not appended to it."""
    monkeypatch.setenv("OCR_ENABLED", "true")
    garbled = MagicMock()
    garbled.get_text.return_value = "ÊÌåÇ ÈÑÇí ÊÓÊ ÇÓÊ " * 5
    document = MagicMock()
    document.__iter__.return_value = [garbled]
    mock_fitz_open.return_value = document
    mock_extract_page_text.return_value = "متن بازخوانی شده"

    sections = parse_document_sections(b"pdf", "broken-font.pdf")

    assert [section.text for section in sections] == ["متن بازخوانی شده"]
    document.close.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.extract_page_text")
@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_mixed_pdf_only_ocrs_scanned_pages(
    mock_fitz_open: MagicMock,
    mock_extract_page_text: MagicMock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Native PDF pages remain the fast path in mixed documents."""
    monkeypatch.setenv("OCR_ENABLED", "true")
    native, scanned = MagicMock(), MagicMock()
    native.get_text.return_value = "native text"
    scanned.get_text.return_value = ""
    document = MagicMock()
    document.__iter__.return_value = [native, scanned]
    mock_fitz_open.return_value = document
    mock_extract_page_text.return_value = "ocr text"

    sections = parse_document_sections(b"pdf", "mixed.pdf")

    assert [(item.text, item.metadata) for item in sections] == [
        ("native text", {"page": 1}),
        ("ocr text", {"page": 2}),
    ]
    mock_extract_page_text.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.extract_page_text")
@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_sparse_pdf_text_layer_does_not_hide_image_evidence(
    mock_open: MagicMock,
    mock_extract: MagicMock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A short searchable title must not suppress OCR of the main page image."""
    monkeypatch.setenv("OCR_ENABLED", "1")
    page = MagicMock()
    page.get_text.return_value = "Invoice 42"
    page.get_images.return_value = [(1,)]
    mock_open.return_value.__iter__.return_value = [page]
    mock_extract.return_value = "Total due: 425 euros"

    sections = parse_document_sections(b"pdf", "invoice.pdf")

    assert sections == [
        ParsedSection("Invoice 42\n\nTotal due: 425 euros", {"page": 1})
    ]
    mock_extract.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.extract_page_text")
@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_pdf_ocr_page_limit_is_enforced(
    mock_fitz_open: MagicMock,
    mock_extract_page_text: MagicMock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """OCR work is bounded independently from the upload byte limit."""
    monkeypatch.setenv("OCR_ENABLED", "1")
    monkeypatch.setenv("OCR_MAX_PAGES", "1")
    pages = [MagicMock(), MagicMock()]
    for page in pages:
        page.get_text.return_value = ""
    document = MagicMock()
    document.__iter__.return_value = pages
    mock_fitz_open.return_value = document
    mock_extract_page_text.return_value = "ocr text"

    parsed = parse_document_file(b"pdf", "large-scan.pdf")

    assert [section.metadata for section in parsed.sections] == [{"page": 1}]
    assert parsed.notices == ("ocr_page_limit",)
    assert mock_extract_page_text.call_count == 1
    document.close.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_password_protected_pdf_is_reported_as_encrypted(
    mock_fitz_open: MagicMock,
) -> None:
    document = MagicMock()
    document.needs_pass = True
    mock_fitz_open.return_value = document

    with pytest.raises(DocumentError) as raised:
        parse_document_file(b"%PDF-1.7", "secret.pdf")

    assert raised.value.code == "encrypted_document"
    document.close.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.extract_image_text")
def test_office_images_past_the_limit_are_skipped_not_fatal(
    mock_extract: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Text around the 31st image is still indexed; a notice says what was skipped."""
    from backend.infrastructure.parsers.document_parser import recognize_images
    from backend.infrastructure.parsers.ocr import OCRError, OCRSettings

    monkeypatch.setenv("OCR_MAX_IMAGES", "1")
    settings = OCRSettings.from_environment()
    notices: list[str] = []
    mock_extract.return_value = "chart caption"

    within_limit = recognize_images([b"img", b"img"], settings, notices)
    mock_extract.side_effect = OCRError("unreadable")
    unreadable = recognize_images(
        [b"img"], OCRSettings(True, "fas", 200, 30, 30), notices
    )

    assert within_limit == ["chart caption", ""]
    assert unreadable == [""]
    assert notices == ["ocr_image_limit", "ocr_partial"]


def test_missing_tesseract_skips_images_with_a_notice() -> None:
    """Word text is indexed even when the images cannot be read."""
    from backend.infrastructure.parsers.document_parser import recognize_images
    from backend.infrastructure.parsers.ocr import OCRSettings, OCRUnavailableError

    notices: list[str] = []
    with patch(
        "backend.infrastructure.parsers.document_parser.extract_image_text",
        side_effect=OCRUnavailableError("missing"),
    ):
        result = recognize_images(
            [b"img"], OCRSettings(True, "fas", 200, 30, 30), notices
        )
    assert result == [""]
    assert notices == ["ocr_unavailable"]


@patch("backend.infrastructure.parsers.document_parser.extract_page_text")
@patch("backend.infrastructure.parsers.document_parser.pymupdf.open")
def test_missing_tesseract_keeps_text_pages_of_a_pdf(
    mock_fitz_open: MagicMock,
    mock_extract_page_text: MagicMock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """One scanned page no longer rejects a PDF whose other pages have text."""
    from backend.infrastructure.parsers.ocr import OCRUnavailableError

    monkeypatch.setenv("OCR_ENABLED", "1")
    text_page, scanned_page = MagicMock(), MagicMock()
    text_page.get_text.return_value = "این صفحه متن کافی دارد و نیازی به OCR ندارد."
    scanned_page.get_text.return_value = ""
    document = MagicMock()
    document.needs_pass = False
    document.__iter__.return_value = [text_page, scanned_page]
    mock_fitz_open.return_value = document
    mock_extract_page_text.side_effect = OCRUnavailableError("missing")

    parsed = parse_document_file(b"pdf", "mixed.pdf")

    assert [section.metadata for section in parsed.sections] == [{"page": 1}]
    assert parsed.notices == ("ocr_unavailable",)

    document.__iter__.return_value = [scanned_page]
    with pytest.raises(DocumentError) as raised:
        parse_document_file(b"pdf", "scan.pdf")
    assert raised.value.code == "ocr_unavailable"


@pytest.mark.parametrize(
    ("payload", "expected"),
    [
        # Windows-1256 has Arabic Yeh (ي) only; ingestion later maps it to ی.
        ("گزارش نهايي پروژه با ي و ک".encode("cp1256"), "گزارش نهايي پروژه با ي و ک"),
        ("سلام دنیا".encode("utf-16"), "سلام دنیا"),
        ("Café résumé naïve".encode("cp1252"), "Café résumé naïve"),
        ("﻿متن UTF-8".encode(), "متن UTF-8"),
    ],
)
def test_legacy_and_unicode_text_encodings_are_read(
    payload: bytes, expected: str
) -> None:
    """Persian Windows (1256), UTF-16, Western Windows, and UTF-8 BOM files."""
    sections = parse_document_sections(payload, "notes.txt")
    assert sections[0].text == expected


def test_non_utf8_text_has_an_encoding_code() -> None:
    with pytest.raises(DocumentError) as raised:
        parse_document_file(b"abc \x81 def ghi", "notes.txt")
    assert raised.value.code == "text_encoding"


def test_chunk_text() -> None:
    # Simulate a Persian text
    persian_text = "این یک متن تستی برای سیستم پردازش زبان طبیعی پارس رگ است. " * 50
    nodes = chunk_text(persian_text, metadata={"source": "test.txt"})

    assert len(nodes) > 0
    assert nodes[0].metadata["source"] == "test.txt"
    # Ensure splitting happened (50 sentences should be split into multiple chunks)
    assert len(nodes) > 1
    assert len(nodes[0].text) > 0


def test_page_bridge_keeps_both_sides_bounded_and_traceable() -> None:
    """A sentence split by a page turn remains one searchable evidence chunk."""
    previous = " ".join(["prefix"] * 150 + ["invoice", "total", "is"])
    current = " ".join(["425", "euros"] + ["suffix"] * 150)

    bridge = bridge_adjacent_pages(
        previous, current, filename="scan.pdf", previous_page=3, current_page=4
    )

    assert bridge is not None
    assert "invoice total is\n[page 4] 425 euros" in bridge.text
    assert bridge.metadata == {
        "filename": "scan.pdf",
        "page": 3,
        "page_end": 4,
        "kind": "page_bridge",
    }
    assert len(bridge.text.split()) <= 184
    assert (
        bridge_adjacent_pages(
            previous, current, filename="scan.pdf", previous_page=3, current_page=5
        )
        is None
    )


def test_real_two_page_pdf_preserves_one_split_fact() -> None:
    """A real PDF page turn keeps both halves available to one retrieval chunk."""
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "The invoice total is")
    document.new_page().insert_text((72, 72), "425 euros due today")
    payload = document.tobytes()
    document.close()

    sections = parse_document_sections(payload, "invoice.pdf")
    bridge = bridge_adjacent_pages(
        sections[0].text,
        sections[1].text,
        filename="invoice.pdf",
        previous_page=sections[0].metadata["page"],
        current_page=sections[1].metadata["page"],
    )

    assert bridge is not None
    assert "invoice total is\n[page 2] 425 euros" in bridge.text


def test_real_docx_parsing() -> None:
    """A Persian Word file with a heading, paragraphs, and a table parses fully."""
    from docx import Document as WordDocument

    document = WordDocument()
    document.add_heading("نیازمندی‌های تست", level=1)
    document.add_paragraph("تست واحد باید ورودی‌های معمولی و نامعتبر را پوشش دهد.")
    document.add_paragraph("تست یکپارچگی جریان داده بین ماژول‌ها را بررسی می‌کند.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "شناسه"
    table.cell(0, 1).text = "ماژول‌ها"
    table.cell(1, 0).text = "IS002"
    table.cell(1, 1).text = "ورود و داشبورد"
    buffer = BytesIO()
    document.save(buffer)

    text = parse_document(buffer.getvalue(), "نیازمندی‌های تست.docx")

    assert "تست واحد باید ورودی‌های معمولی" in text
    assert "جریان داده بین ماژول‌ها" in text
    assert "IS002" in text and "ورود و داشبورد" in text


@pytest.mark.parametrize(
    ("filename", "payload", "expected"),
    [
        ("notes.txt", "متن ساده".encode(), "متن ساده"),
        ("guide.md", "# راهنما\n\nجزئیات".encode(), "# راهنما"),
        (
            "data.json",
            json.dumps({"name": "پارس‌رگ"}, ensure_ascii=False).encode(),
            "پارس‌رگ",
        ),
        ("table.csv", b"name,value\nalpha,42", "alpha"),
        ("page.html", b"<h1>Title</h1><p>Body</p>", "Title"),
    ],
)
def test_textual_formats_are_parsed(
    filename: str, payload: bytes, expected: str
) -> None:
    """Common UTF-8 knowledge formats are accepted without lossy conversion."""
    sections = parse_document_sections(payload, filename)

    assert expected in sections[0].text
    assert sections[0].metadata == {"section": 1}


@patch("backend.infrastructure.parsers.document_parser.extract_image_text")
def test_standalone_image_uses_ocr(
    mock_extract: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A raster image is a first-class OCR document."""
    monkeypatch.setenv("OCR_ENABLED", "1")
    mock_extract.return_value = "متن روی تصویر"

    sections = parse_document_sections(b"image", "scan.png")

    assert sections == [ParsedSection("متن روی تصویر", {"page": 1})]


@patch("backend.infrastructure.parsers.document_parser.extract_image_text")
def test_docx_embedded_image_is_ocred_when_enabled(
    mock_extract: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Pictures embedded in Word contribute OCR text at their paragraph location."""
    monkeypatch.setenv("OCR_ENABLED", "1")
    mock_extract.return_value = "نوشته داخل تصویر ورد"
    image = BytesIO()
    Image.new("RGB", (400, 200), "white").save(image, format="PNG")
    document = Document()
    document.add_picture(BytesIO(image.getvalue()))
    payload = BytesIO()
    document.save(payload)

    sections = parse_document_sections(payload.getvalue(), "image-only.docx")

    assert sections[0].text == "[تصویر 1]\nنوشته داخل تصویر ورد"
    assert sections[0].metadata == {"paragraph": 1, "image": 1}
    mock_extract.assert_called_once()


@patch("backend.infrastructure.parsers.document_parser.extract_image_text")
def test_docx_images_are_cited_with_number_caption_and_paragraph(
    mock_extract: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Each picture is its own section: numbered, captioned, placed after its paragraph."""
    monkeypatch.setenv("OCR_ENABLED", "1")
    mock_extract.side_effect = ["متن تصویر اول", "", "متن تصویر سوم"]
    image = BytesIO()
    Image.new("RGB", (400, 200), "white").save(image, format="PNG")
    document = Document()
    document.add_heading("نتایج آزمون", level=1)
    document.add_paragraph("پاراگراف پیش از تصویر.")
    document.add_picture(BytesIO(image.getvalue()))
    document.add_paragraph("شکل ۳-۱: صفحه اجرای آزمون")
    document.add_picture(BytesIO(image.getvalue()))
    document.add_paragraph("متن پایانی بدون شرح.")
    document.add_picture(BytesIO(image.getvalue()))
    payload = BytesIO()
    document.save(payload)

    sections = parse_document_sections(payload.getvalue(), "report.docx")

    images = [s for s in sections if "image" in s.metadata]
    assert [s.metadata for s in images] == [
        {"paragraph": 2, "image": 1},
        {"paragraph": 4, "image": 3},
    ]
    assert (
        images[0].text
        == "نتایج آزمون\n[تصویر 1] شکل ۳-۱: صفحه اجرای آزمون\nپاراگراف پیش از تصویر.\nمتن تصویر اول"
    )
    assert (
        images[1].text == "نتایج آزمون\n[تصویر 3]\nمتن پایانی بدون شرح.\nمتن تصویر سوم"
    )
    text = [s for s in sections if "image" not in s.metadata]
    assert text[0].metadata == {"paragraph": 2, "paragraph_end": 4}
    assert "متن تصویر" not in text[0].text


@patch("backend.infrastructure.parsers.document_parser.extract_image_text")
def test_pptx_picture_is_ocred_with_slide_metadata(
    mock_extract: MagicMock, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Pictures embedded in PowerPoint keep their slide citation after OCR."""
    monkeypatch.setenv("OCR_ENABLED", "1")
    mock_extract.return_value = "نوشته داخل تصویر پاورپوینت"
    image = BytesIO()
    Image.new("RGB", (400, 200), "white").save(image, format="PNG")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(image.getvalue()), Inches(1), Inches(1))
    payload = BytesIO()
    presentation.save(payload)

    sections = parse_document_sections(payload.getvalue(), "image-only.pptx")

    assert sections[0].text == "[تصویر 1]\nنوشته داخل تصویر پاورپوینت"
    assert sections[0].metadata == {"slide": 1, "image": 1}
    mock_extract.assert_called_once()


def test_parallel_ocr_keeps_document_order(monkeypatch: pytest.MonkeyPatch) -> None:
    """Images finish in any order but their text stays in document order."""
    import time

    from backend.infrastructure.parsers.document_parser import recognize_images
    from backend.infrastructure.parsers.ocr import OCRSettings

    monkeypatch.setenv("OCR_WORKERS", "4")
    seen_single_thread: list[bool] = []

    def slow_read(image: bytes, settings: OCRSettings) -> str:
        seen_single_thread.append(settings.single_thread)
        time.sleep(0.05 * (4 - int(image)))
        return f"text {image.decode()}"

    with patch(
        "backend.infrastructure.parsers.document_parser.extract_image_text",
        side_effect=slow_read,
    ):
        texts = recognize_images(
            [b"1", b"2", b"3"], OCRSettings(True, "fas", 200, 30, 30), []
        )

    assert texts == ["text 1", "text 2", "text 3"]
    assert all(seen_single_thread)


@pytest.mark.parametrize(("value", "expected"), [("1", 1), ("3", 3), ("x", None)])
def test_ocr_workers_are_configurable(
    monkeypatch: pytest.MonkeyPatch, value: str, expected: int | None
) -> None:
    from backend.infrastructure.parsers.document_parser import ocr_workers

    monkeypatch.setenv("OCR_WORKERS", value)
    assert ocr_workers() == (expected if expected is not None else ocr_workers())
    assert 1 <= ocr_workers() <= 16


def test_image_previews_are_small_webp_and_skip_unreadable_data() -> None:
    import io

    from PIL import Image

    from backend.infrastructure.parsers.document_parser import image_preview

    source = io.BytesIO()
    Image.new("RGB", (2400, 1200), "white").save(source, format="PNG")
    preview = image_preview(source.getvalue())

    assert preview is not None and preview[:4] == b"RIFF" and preview[8:12] == b"WEBP"
    with Image.open(io.BytesIO(preview)) as image:
        assert max(image.size) == 1280
    assert image_preview(b"not an image") is None


def test_only_images_that_became_sections_get_previews(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import io

    from PIL import Image

    from backend.core.domain.documents import ParsedSection
    from backend.infrastructure.parsers.document_parser import _section_previews

    png = io.BytesIO()
    Image.new("RGB", (40, 40), "red").save(png, format="PNG")
    sections = [
        ParsedSection("text", {"paragraph": 1}),
        ParsedSection("[تصویر 2] chart", {"paragraph": 1, "image": 2}),
    ]
    blobs = {1: png.getvalue(), 2: png.getvalue()}

    assert list(_section_previews(sections, blobs)) == [2]
    monkeypatch.setenv("SOURCE_IMAGE_PREVIEWS", "0")
    assert _section_previews(sections, blobs) == {}


def test_automatic_ocr_workers_use_the_available_cpus() -> None:
    from backend.infrastructure.parsers.document_parser import default_ocr_workers

    affinity = "backend.infrastructure.parsers.document_parser.os.sched_getaffinity"
    with patch(affinity, create=True, return_value={0, 1, 2}):
        assert default_ocr_workers() == 3
    with patch(affinity, create=True, return_value=set(range(20))):
        assert default_ocr_workers() == 8
