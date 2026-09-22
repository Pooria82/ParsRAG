import io
import json
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import Any

import fitz  # type: ignore  # PyMuPDF
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.core.exceptions import EmptyDocumentError
from backend.core.upload_policy import IMAGE_EXTENSIONS, TEXT_EXTENSIONS
from backend.infrastructure.parsers.ocr import (
    OCRError,
    OCRImageLimitError,
    OCRPageLimitError,
    OCRSettings,
    extract_image_text,
    extract_page_text,
)

__all__ = [
    "EmptyDocumentError",
    "ParsedSection",
    "parse_document",
    "parse_document_sections",
]


@dataclass(frozen=True)
class ParsedSection:
    """Text extracted from a traceable document location."""

    text: str
    metadata: dict[str, int]


class _VisibleTextExtractor(HTMLParser):
    """Collect visible text from simple HTML without executing or retaining markup."""

    def __init__(self) -> None:
        """Initialize an empty text collector."""
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        """Ignore executable/style blocks and separate structural elements."""
        del attrs
        if tag in {"script", "style"}:
            self._ignored_depth += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        """Restore collection after ignored blocks and preserve block boundaries."""
        if tag in {"script", "style"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        """Collect text outside ignored elements."""
        if not self._ignored_depth:
            self.parts.append(data)


def parse_document(file_bytes: bytes, filename: str) -> str:
    """Extract text content from one supported knowledge-file byte stream.

    Preserves chronological document structure and converts tables into
    well-formatted Markdown with row-level header associations.

    Args:
        file_bytes (bytes): The raw bytes of the file.
        filename (str): The name of the file to determine its type.

    Returns:
        str: The extracted text from the document.

    Raises:
        EmptyDocumentError: If the document contains no text.
        ValueError: If the file format is unsupported.
    """
    sections = parse_document_sections(file_bytes, filename)
    if filename.lower().endswith(".pptx"):
        return "\n\n".join(
            f"--- اسلاید {section.metadata['slide']} ---\n\n{section.text}"
            for section in sections
        )
    return "\n\n".join(section.text for section in sections)


def parse_document_sections(file_bytes: bytes, filename: str) -> list[ParsedSection]:
    """Extracts text with page, slide, paragraph, or section metadata."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        sections = _parse_pdf_sections(file_bytes)
    elif lower_name.endswith(".docx"):
        sections = _parse_docx_sections(file_bytes)
    elif lower_name.endswith(".pptx"):
        sections = _parse_pptx_sections(file_bytes)
    elif lower_name.endswith(IMAGE_EXTENSIONS):
        sections = _parse_image_sections(file_bytes)
    elif lower_name.endswith(TEXT_EXTENSIONS):
        sections = _parse_text_sections(file_bytes, lower_name)
    else:
        raise ValueError("Unsupported file format.")
    if not sections:
        if lower_name.endswith(".pdf"):
            raise EmptyDocumentError("Scanned PDFs require OCR, but OCR is disabled.")
        if lower_name.endswith(IMAGE_EXTENSIONS):
            raise EmptyDocumentError("Image-based documents require enabled OCR.")
        raise EmptyDocumentError("The document contains no text.")
    return sections


def _parse_docx_sections(file_bytes: bytes) -> list[ParsedSection]:
    """Extract traceable paragraphs and tables from DOCX bytes."""
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError("Corrupted or invalid DOCX document.") from exc
    settings = OCRSettings.from_environment()
    sections: list[ParsedSection] = []
    paragraph = 0
    ocr_images = 0
    for section_index, child in enumerate(
        document.element.body.iterchildren(), start=1
    ):
        if isinstance(child, CT_P):
            paragraph_object = Paragraph(child, document)
            parts = (
                [paragraph_object.text.strip()] if paragraph_object.text.strip() else []
            )
            if settings.enabled:
                for image in _docx_paragraph_images(paragraph_object, document):
                    text, ocr_images = _extract_embedded_image(
                        image, settings, ocr_images, "Word"
                    )
                    if text:
                        parts.append(text)
            value = "\n\n".join(parts).strip()
            if value:
                paragraph += 1
                sections.append(ParsedSection(value, {"paragraph": paragraph}))
        elif isinstance(child, CT_Tbl):
            value = _format_docx_table(Table(child, document)).strip()
            if value:
                sections.append(ParsedSection(value, {"section": section_index}))
    return sections


def _parse_pptx_sections(file_bytes: bytes) -> list[ParsedSection]:
    """Extract traceable slide content and tables from PPTX bytes."""
    try:
        presentation = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ValueError("Corrupted or invalid PPTX document.") from exc
    settings = OCRSettings.from_environment()
    sections: list[ParsedSection] = []
    ocr_images = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                parts.append(_format_pptx_table(shape.table))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and settings.enabled:
                text, ocr_images = _extract_embedded_image(
                    shape.image.blob, settings, ocr_images, "PowerPoint"
                )
                if text:
                    parts.append(text)
            elif hasattr(shape, "text") and shape.text and shape.text.strip():
                parts.append(shape.text.strip())
        if text := "\n\n".join(part for part in parts if part).strip():
            sections.append(ParsedSection(text, {"slide": slide_index}))
    return sections


def _docx_paragraph_images(paragraph: Paragraph, document: Any) -> list[bytes]:
    """Return image blobs referenced by one Word paragraph in document order."""
    images: list[bytes] = []
    relationship_ids = paragraph._p.xpath(".//a:blip/@r:embed")
    for relationship_id in relationship_ids:
        related_part = document.part.related_parts.get(relationship_id)
        blob = getattr(related_part, "blob", None)
        if isinstance(blob, bytes):
            images.append(blob)
    return images


def _extract_embedded_image(
    image: bytes, settings: OCRSettings, processed: int, source: str
) -> tuple[str, int]:
    """OCR one bounded Office image and return its text and updated counter."""
    next_count = processed + 1
    if next_count > settings.max_images:
        raise OCRImageLimitError(
            f"Documents are limited to {settings.max_images} OCR images."
        )
    try:
        return extract_image_text(image, settings), next_count
    except OCRError as exc:
        raise ValueError(
            f"An embedded {source} image could not be processed: {exc}"
        ) from exc


def _parse_image_sections(file_bytes: bytes) -> list[ParsedSection]:
    """Extract OCR text from a standalone raster image."""
    settings = OCRSettings.from_environment()
    if not settings.enabled:
        return []
    try:
        text = extract_image_text(file_bytes, settings)
    except OCRError as exc:
        raise ValueError(f"The image could not be processed: {exc}") from exc
    return [ParsedSection(text, {"page": 1})] if text else []


def _decode_text(file_bytes: bytes) -> str:
    """Decode bounded text formats as UTF-8 while rejecting binary payloads."""
    if b"\x00" in file_bytes:
        raise ValueError("The text file contains binary data.")
    try:
        return file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Text files must use UTF-8 encoding.") from exc


def _parse_text_sections(file_bytes: bytes, filename: str) -> list[ParsedSection]:
    """Parse structured and plain UTF-8 knowledge formats."""
    text = _decode_text(file_bytes)
    if filename.endswith(".json"):
        try:
            value = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ValueError("The JSON document is invalid.") from exc
        text = json.dumps(value, ensure_ascii=False, indent=2)
    elif filename.endswith((".html", ".htm", ".xml")):
        parser = _VisibleTextExtractor()
        parser.feed(text)
        text = "\n".join(
            line.strip() for line in "".join(parser.parts).splitlines() if line.strip()
        )
    text = text.strip()
    return [ParsedSection(text, {"section": 1})] if text else []


def _clean_cell_text(text: str) -> str:
    """Cleans cell text by normalizing whitespace and escaping markdown pipes."""
    clean = " ".join(text.split())
    return clean.replace("|", "\\|")


def _format_table_grid(grid_rows: list[list[str]]) -> str:
    """Converts a grid of strings into Markdown table with structured row records.

    Args:
        grid_rows (list[list[str]]): 2D list of cleaned cell strings.

    Returns:
        str: Formatted Markdown table and structured records.
    """
    if not grid_rows:
        return ""

    num_rows = len(grid_rows)
    max_cols = max(len(r) for r in grid_rows)

    # 1x1 callout or code box
    if num_rows == 1 and max_cols == 1:
        content = grid_rows[0][0]
        return f"> [کادر محتوا]:\n> {content}"

    # Pad any short rows to max_cols
    for row in grid_rows:
        while len(row) < max_cols:
            row.append("-")

    # Header resolution
    headers = [
        col if col else f"ستون_{idx + 1}" for idx, col in enumerate(grid_rows[0])
    ]

    # Build Markdown table
    md_lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join([":---"] * max_cols) + " |",
    ]
    for row in grid_rows[1:]:
        row_cells = [c if c else "-" for c in row]
        md_lines.append("| " + " | ".join(row_cells) + " |")

    table_md = "\n".join(md_lines)

    # Build Structured Row Records for tables with >= 2 columns and >= 2 rows
    # This prevents orphaned cell values when chunks split across table rows
    if max_cols >= 2 and num_rows >= 2:
        record_lines = ["\n[سوابق ردیف‌های جدول]:"]
        for r_idx, row in enumerate(grid_rows[1:], start=1):
            items = []
            for h, val in zip(headers, row, strict=False):
                cell_val = val if val else "-"
                items.append(f"[{h}]: {cell_val}")
            record_lines.append(f"- سطر {r_idx}: " + " | ".join(items))
        table_md += "\n" + "\n".join(record_lines)

    return table_md


def _format_docx_table(table: Table) -> str:
    """Formats a DOCX table into Markdown and row records, handling merged cells."""
    if not table.rows:
        return ""

    grid_rows: list[list[str]] = []
    for row in table.rows:
        # Deduplicate horizontally merged cells
        unique_cells: list[_Cell] = []
        for cell in row.cells:
            if not unique_cells or cell._tc != unique_cells[-1]._tc:
                unique_cells.append(cell)
        row_vals = [_clean_cell_text(cell.text) for cell in unique_cells]
        if any(row_vals):
            grid_rows.append(row_vals)

    return _format_table_grid(grid_rows)


def _format_pptx_table(table: Any) -> str:
    """Formats a PPTX table shape into Markdown and row records."""
    grid_rows: list[list[str]] = []
    for row in table.rows:
        row_vals = [_clean_cell_text(cell.text) for cell in row.cells]
        if any(row_vals):
            grid_rows.append(row_vals)

    return _format_table_grid(grid_rows)


def _parse_pdf_sections(file_bytes: bytes) -> list[ParsedSection]:
    """Extracts ordered PDF pages, using OCR only for scanned pages."""
    try:
        document = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise ValueError("Corrupted or invalid PDF document.") from exc

    settings = OCRSettings.from_environment()
    sections: list[ParsedSection] = []
    scanned_pages = 0
    try:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if not text:
                scanned_pages += 1
                if not settings.enabled:
                    continue
                if scanned_pages > settings.max_pages:
                    raise OCRPageLimitError(
                        f"Scanned PDFs are limited to {settings.max_pages} OCR pages."
                    )
                try:
                    text = extract_page_text(page, settings)
                except OCRError as exc:
                    raise ValueError(
                        f"Scanned PDFs could not be processed: {exc}"
                    ) from exc
            if text:
                sections.append(
                    ParsedSection(text=text, metadata={"page": page_number})
                )
    finally:
        document.close()
    return sections
